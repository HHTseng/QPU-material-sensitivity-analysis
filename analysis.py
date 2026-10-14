"""Explicit weighted statistics and comparisons of complete simulation arrays."""

import math
import copy
import inspect
import json
from pathlib import Path

import physics
import records

ROOT = Path(__file__).resolve().parent

import numpy as np


def prepare_counts(blocks, design=None):
    if isinstance(blocks,dict) and 'groups' in blocks and 'totals' in blocks: return blocks
    require_electrodes=False
    labels = tuple(str(v) for v in design["position_strata"])
    weights = dict(design["stratum_weights"])
    replicas = tuple(design["replica_ids"])
    ew = np.asarray(design["electrode_weights"], dtype=float)
    if not labels or not replicas or len(set(replicas)) != len(replicas):
        raise ValueError("positions and unique replica IDs must be nonempty")
    if not design.get("design_hash"):
        raise ValueError("design_hash is required")
    if design.get("injection_law", "uniform_surface") != "uniform_surface":
        raise ValueError("area weighting requires uniform_surface injection")
    if set(labels) != set(weights) or not all(math.isfinite(v) and v > 0 for v in weights.values()):
        raise ValueError("position labels must match the positive finite area weights")
    if not math.isclose(sum(weights.values()), 1., rel_tol=0., abs_tol=1e-12):
        raise ValueError("area weights must sum to one")
    if ew.ndim != 1 or not ew.size or not np.isfinite(ew).all() or np.any(ew < 0) or not np.isclose(ew.sum(), 1., rtol=0., atol=1e-12):
        raise ValueError("electrode weights must be non-negative and sum to one")
    events, energy = design["events_per_task"], design["gun_energy_eV"]
    if not isinstance(events, (int, np.integer)) or events <= 0 or not math.isfinite(energy) or energy <= 0:
        raise ValueError("positive integer events and finite positive excitation required")
    expected = {(p, r) for p in range(len(labels)) for r in replicas}
    indexed = {}
    for block in blocks:
        key = (int(block["position"]), int(block["replica"]))
        if key in indexed or key not in expected:
            raise ValueError(f"duplicate or undeclared block {key}")
        if block["events"] != events:
            raise ValueError(f"block {key} has {block['events']} events, expected {events}")
        total = block["total_qps"]
        if not math.isfinite(total) or total < 0:
            raise ValueError(f"invalid QP total in block {key}")
        electrodes = block.get("per_electrode_qps")
        if electrodes is None:
            if require_electrodes:
                raise ValueError(f"block {key} lacks electrode counts")
        else:
            a = np.asarray(electrodes, dtype=float)
            if a.shape != ew.shape or not np.isfinite(a).all() or np.any(a < 0):
                raise ValueError(f"invalid electrode counts in block {key}")
            if not math.isclose(math.fsum(a), total, rel_tol=0., abs_tol=1e-9):
                raise ValueError(f"electrode counts disagree with total in block {key}")
        if block.get("n_hits") is not None and block["n_hits"] < 0:
            raise ValueError(f"negative hit count in block {key}")
        indexed[key] = block
    if set(indexed) != expected:
        raise ValueError(f"incomplete block set: missing {len(expected - set(indexed))} of {len(expected)}")
    totals=np.asarray([[indexed[(i,r)]['total_qps'] for r in replicas] for i in range(len(labels))])
    counts=None; weighted=None; worst=None
    if all(block.get('per_electrode_qps') is not None for block in indexed.values()):
        counts=np.asarray([[indexed[(i,r)]['per_electrode_qps'] for r in replicas] for i in range(len(labels))])
        weighted=np.asarray([[float(row@ew)/(events*energy) for row in site] for site in counts])
        worst=counts.max(axis=2)/(events*energy)
    grouped={}
    for i,label in enumerate(labels): grouped.setdefault(label,[]).append(i)
    return {'design':design,'totals':totals,'counts':counts,'weighted':weighted,'worst':worst,
            'groups':[(label,positions,weights[label]) for label,positions in grouped.items()]}


def _tail_mean(pairs, probability=.05):
    pairs = sorted(pairs, reverse=True)
    cutoff = probability * sum(w for _, w in pairs)
    numerator = mass = 0.
    for value, weight in pairs:
        take = min(weight, cutoff - mass)
        if take <= 0:
            break
        numerator += value * take
        mass += take
    if mass <= 0:
        raise ValueError("tail mean needs positive area")
    return float(numerator / mass)


def _resample(regions, draws, seed, tail=False):
    if not isinstance(draws, int) or draws < 2:
        raise ValueError("at least two integer resampling draws required")
    rng = np.random.default_rng(seed)
    estimates = np.empty(draws)
    for draw in range(draws):
        total, sampled = 0., []
        for _, positions, weight, matrix in regions:
            selected = rng.choice(len(positions), size=len(positions), replace=True)
            selected_values = matrix[selected]
            replica_index = rng.choice(selected_values.shape[1], size=selected_values.shape, replace=True)
            replica_means = np.take_along_axis(selected_values, replica_index, axis=1).mean(axis=1)
            if tail:
                sampled.extend((value, weight / len(positions)) for value in replica_means)
            else:
                total += weight * float(replica_means.mean())
        estimates[draw] = _tail_mean(sampled) if tail else total
    return float(np.std(estimates, ddof=1))


def paired_difference(blocks_a, blocks_b, design, *, bootstrap_draws=1000, bootstrap_seed=23,
                      objective='device_junction_qps_per_eV'):
    """A − B with matched sites, replicas and declared weights before resampling."""
    a=prepare_counts(blocks_a,design); b=prepare_counts(blocks_b,design)
    historical=objective in ('total_qps_per_primary','total_QPs_per_primary','junction_qps_per_primary')
    if historical:
        av=a['totals']/design['events_per_task']; bv=b['totals']/design['events_per_task']
        groups=[('surface',list(range(len(av))),1.)]
    else:
        if objective not in ('device_junction_qps_per_eV','device_weighted_junction_qps_per_energy') or a['weighted'] is None or b['weighted'] is None:
            raise ValueError('paired objective requires complete electrode counts')
        av=a['weighted']; bv=b['weighted']; groups=a['groups']
    delta=av-bv; b=bv
    regions=[(label,positions,weight,delta[positions]) for label,positions,weight in groups]
    value = sum(weight * float(np.mean([float(np.mean(delta[p])) for p in positions]))
                for _, positions, weight, _ in regions)
    if historical:
        se = float(np.sqrt(delta.var(axis=1, ddof=1).sum() / delta.shape[1]) / len(delta)) if delta.shape[1] >= 2 else None
    else:
        se = _resample(regions, bootstrap_draws, bootstrap_seed)
    baseline = sum(weight * float(np.mean([float(np.mean(b[p])) for p in positions]))
                   for _, positions, weight, _ in regions)
    return {"value": float(value), "standard_error": se,
            "relative_to_b": float(value / baseline) if baseline else None,
            "z": float(value / se) if se else None,
            "confidence_interval_95": [float(value - 1.96 * se), float(value + 1.96 * se)] if se is not None else None,
            "n_sites": len(delta), "sites_favouring_a": int(np.sum(delta.mean(axis=1) < 0)),
            'units': 'summed_junction_QPs_per_primary' if historical else 'weighted_junction_QPs_per_injected_eV',
            'uncertainty_method': 'within-site paired replica variance' if historical else 'hierarchical site and replica resampling'}


def spatial_convergence(coarse, fine, *, baseline_coarse=None, baseline_fine=None,
                        relative_limit=.05, sigma_limit=2., leverage_limit=.10):
    """Declared area-convergence checks; region shifts identify sites to refine."""
    a, b = float(coarse["value"]), float(fine["value"])
    if a <= 0 or b <= 0:
        raise ValueError("positive measured yields required for relative convergence")
    shift = abs(b - a) / b
    combined = math.hypot(coarse["standard_error"], fine["standard_error"])
    region_shift = {}
    if set(coarse["means_by_stratum"]) != set(fine["means_by_stratum"]):
        raise ValueError("convergence comparison requires the same physical regions")
    for name, value in fine["means_by_stratum"].items():
        old = coarse["means_by_stratum"][name]
        region_shift[name] = abs(value - old) / abs(value) if value else (0. if old == 0 else float("inf"))
    reduction_shift = None
    if baseline_coarse is not None and baseline_fine is not None:
        if baseline_coarse <= 0 or baseline_fine <= 0:
            raise ValueError("baseline yields must be positive")
        reduction_shift = abs(a / baseline_coarse - b / baseline_fine)
    checks = {"relative_shift": shift <= relative_limit,
              "within_combined_error": combined > 0 and abs(b - a) <= sigma_limit * combined,
              "reduction_stability": reduction_shift is None or reduction_shift <= relative_limit,
              "maximum_site_leverage": fine["maximum_site_leverage"] <= leverage_limit,
              "every_region": all(v <= relative_limit for v in region_shift.values())}
    return {"passed": all(checks.values()), "checks": checks, "relative_shift": shift,
            "combined_standard_error": combined, "reduction_shift": reduction_shift,
            "region_relative_shifts": region_shift,
            "refine_regions": [k for k, v in region_shift.items() if v > relative_limit],
            "ranking_checked": False}


def sensitivity(samples, outcomes, names=None, *, problem=None, method="pearson", log_outcome=False, seed=7, num_levels=4):
    """Pearson correlations or Morris μ, μ*, σ; failed outcomes remain absent."""
    X, y = np.asarray(samples, dtype=float), np.asarray(outcomes, dtype=float)
    if X.ndim != 2 or y.shape != (X.shape[0],) or not np.isfinite(X).all() or not np.isfinite(y).all():
        raise ValueError("finite sample matrix and matching complete outcome vector required")
    names = list(names or [f"x{i}" for i in range(X.shape[1])])
    if len(names) != X.shape[1]:
        raise ValueError("parameter names do not match sample columns")
    if log_outcome:
        if np.any(y <= 0):
            raise ValueError("log outcomes must be positive; measured zeros cannot be discarded silently")
        y = np.log10(y)
    if method == "morris":
        if problem is None:
            raise ValueError("Morris analysis needs the original parameter definition")
        from SALib.analyze.morris import analyze
        if len(y) % (X.shape[1] + 1):
            raise ValueError('Morris analysis requires complete trajectories of d+1 samples')
        result = analyze(problem, X, y, seed=seed, num_levels=num_levels, print_to_console=False)
        return {"method": "morris", "names": names, **{k: np.asarray(result[k]).tolist() for k in ("mu", "mu_star", "sigma", "mu_star_conf")}}
    if method != "pearson" or X.shape[0] < 3:
        raise ValueError("Pearson analysis needs at least three samples; unknown method")
    with np.errstate(divide="ignore", invalid="ignore"):
        correlation = np.corrcoef(np.column_stack((X, y)), rowvar=False)[-1, :-1]
    return {"method": "pearson", "names": names, "correlation": [float(c) if np.isfinite(c) else None for c in correlation],
            "constant_parameters": [n for n, c in zip(names, correlation) if not np.isfinite(c)],
            "n_samples": len(y), "log_outcome": log_outcome}


def rank_stability(a, b):
    """Compare matching candidates; flag inversions only when both are resolved."""
    if set(a) != set(b) or len(a) < 2:
        raise ValueError("at least two matching candidates required")
    from scipy.stats import spearmanr, kendalltau
    names = sorted(a)
    av, bv = [a[k]["value"] for k in names], [b[k]["value"] for k in names]
    inversions = []
    for i, first in enumerate(names):
        for second in names[i + 1:]:
            da, db = a[first]["value"] - a[second]["value"], b[first]["value"] - b[second]["value"]
            ea = 1.96 * math.hypot(a[first]["standard_error"], a[second]["standard_error"])
            eb = 1.96 * math.hypot(b[first]["standard_error"], b[second]["standard_error"])
            if da * db < 0 and abs(da) > ea and abs(db) > eb:
                inversions.append([first, second])
    return {"spearman": float(spearmanr(av, bv).statistic), "kendall": float(kendalltau(av, bv).statistic),
            "resolved_inversions": inversions, "stable_within_uncertainty": not inversions}


def curve_distance(sim_t, sim_dg, x_data, y_data, dy_data=None, normalize='none'):
    """Nearest-time residual sum for an explicitly matched measured electrode."""
    t, rate, x, y = [np.asarray(v, dtype=float) for v in (sim_t, sim_dg, x_data, y_data)]
    if t.ndim != 1 or not t.size or rate.shape != t.shape or x.ndim != 1 or y.shape != x.shape or any(not np.isfinite(v).all() for v in (t, rate, x, y)):
        raise ValueError('finite matching simulation and measured time-series required')
    if normalize not in ('none', 'y', 'dy2'):
        raise ValueError('curve normalization must be none, y, or dy2')
    indices = np.array([np.argmin(np.abs(value - t)) for value in x], dtype=int)
    squared = (rate[indices] - y) ** 2
    if normalize == 'y':
        good = np.abs(y) > 1e-12
        return float(np.sum(squared[good] / y[good]))
    if normalize == 'dy2':
        dy = np.asarray(dy_data, dtype=float)
        if dy.shape != y.shape or not np.isfinite(dy).all():
            raise ValueError('finite matching measured uncertainties required')
        good = np.abs(dy) > 1e-12
        return float(np.sum(squared[good] / dy[good] ** 2))
    return float(squared.sum())


def project_materials(target, candidates, feature_spec, weights=None, top=10):
    """Order measured candidates by coverage, then distance within that coverage."""
    weights=weights or {}; scales={}; rows=[]
    for name,mode in feature_spec.items():
        if mode not in ('log','linear'): raise ValueError(f'unknown feature scale {mode!r}')
        values=[float(c['features'][name]) for c in candidates if c['features'].get(name) is not None and np.isfinite(c['features'][name]) and (mode!='log' or c['features'][name]>0)]
        values=np.log(values) if mode=='log' else np.asarray(values)
        mad=float(np.median(np.abs(values-np.median(values))))*1.4826 if len(values)>=2 else 0.
        scales[name]=mad if mad>0 else (float(np.std(values)) or 1.) if len(values)>=2 else 1.
    for candidate in candidates:
        num=den=0.; missing=[]; z_values={}
        for name,mode in feature_spec.items():
            t,c=target.get(name),candidate['features'].get(name)
            if t is None or c is None or not np.isfinite(t) or not np.isfinite(c) or (mode=='log' and (t<=0 or c<=0)):
                missing.append(name); continue
            weight,scale=float(weights.get(name,1.)),scales[name]
            if not math.isfinite(weight) or weight<0 or not math.isfinite(scale) or scale<=0:
                raise ValueError('feature weights must be non-negative and scales positive')
            z=(math.log(c)-math.log(t))/scale if mode=='log' else (float(c)-float(t))/scale
            num+=weight*z*z; den+=weight; z_values[name]=round(z,4)
        if not z_values or den==0: continue
        matched,total=len(z_values),len(feature_spec)
        rows.append({'distance':math.sqrt(num/den),'matched':matched,'total':total,'unmatched_features':missing,
                     'z':z_values,'coverage':matched/total,'stratum':f'{matched}/{total}','id':candidate['id'],'candidate':candidate})
    rows.sort(key=lambda row:(-row['matched'],row['distance']))
    return {'rows':rows[:top],'scales':scales,'missing_features':{name:sum(name in row['unmatched_features'] for row in rows) for name in feature_spec},
            'n_candidates':len(rows),'fully_covered':sum(row['matched']==row['total'] for row in rows)}


def tolerance_summary(centre, moves, *, paired=None, coordinates=None, names=None,
                      improvement_threshold=.02, bound_margin=.02):
    """Local one-variable tests; unavailable evidence cannot establish convergence."""
    value, se = centre["value"], centre["standard_error"]
    if value <= 0:
        raise ValueError("positive centre yield required")
    paired, summaries, improving = paired or {}, {}, []
    measured_error = se is not None and np.isfinite(se) and se >= 0
    variables = sorted({label.rsplit(":", 1)[0] for label in moves})
    complete = bool(measured_error)
    for name in variables:
        labels = [name + ":minus", name + ":plus"]
        tested = [moves[k]["value"] for k in labels if moves.get(k, {}).get("value") is not None]
        swing = max((abs(v - value) / value for v in tested), default=None)
        summaries[name] = {"maximum_relative_swing": swing,
                           "resolved": swing > 4 * se / value if swing is not None and measured_error else None,
                           "below_two_standard_errors": swing < 2 * se / value if swing is not None and measured_error else None}
        for label in labels:
            p = paired.get(label)
            if p is None or p.get("relative_to_b") is None or p.get("z") is None:
                complete = False
                continue
            if p["relative_to_b"] <= -improvement_threshold and p["z"] <= -2.:
                improving.append({"move": label, "relative_change": p["relative_to_b"], "z": p["z"]})
    at_bound = []
    if coordinates is not None:
        at_bound = [name for name, u in zip(names or range(len(coordinates)), coordinates) if u <= bound_margin or u >= 1 - bound_margin]
    return {"variables": summaries, "materially_improving_moves": improving, "variables_near_a_bound": at_bound,
            "complete_paired_tests": complete and bool(variables),
            "locally_converged": bool(variables) and complete and not improving and not at_bound and coordinates is not None,
            "untested_criteria": ["stability after successive model updates", "stability across optimizer random seeds", "variable interactions"]}


def pareto_front(rows, keys, lower_is_better=None):
    lower = set(keys if lower_is_better is None else lower_is_better)
    usable = [r for r in rows if all(r.get(k) is not None and not np.isnan(r[k]) for k in keys)]
    values = np.asarray([[float(row[k]) * (1 if k in lower else -1) for k in keys] for row in usable])
    return [row for i, row in enumerate(usable) if not np.any(np.all(values <= values[i], axis=1) & np.any(values < values[i], axis=1))]


def compare_repeated(a, b, *, confidence=.95, relative_margin=.05, paired=False):
    """Student-t interval for matched or independent outcomes; equivalence needs a margin."""
    from scipy.stats import t
    a, b = np.asarray(a, dtype=float), np.asarray(b, dtype=float)
    if a.ndim != 1 or b.ndim != 1 or min(len(a), len(b)) < 2 or not np.isfinite(a).all() or not np.isfinite(b).all():
        raise ValueError("at least two finite independent outcomes per implementation required")
    if not 0 < confidence < 1 or relative_margin <= 0:
        raise ValueError("confidence must be in (0,1) and equivalence margin positive")
    if paired:
        if a.shape != b.shape:
            raise ValueError('matched realization arrays must have equal lengths')
        se = float((a - b).std(ddof=1) / math.sqrt(len(a)))
        df = len(a) - 1
    else:
        va, vb = a.var(ddof=1) / len(a), b.var(ddof=1) / len(b)
        se = math.sqrt(va + vb)
        df = (va + vb) ** 2 / (va * va / (len(a) - 1) + vb * vb / (len(b) - 1)) if se else float("inf")
    mean_a, mean_b = float(a.mean()), float(b.mean())
    delta = mean_a - mean_b
    radius = float(t.ppf((1 + confidence) / 2, df) * se)
    interval = [delta - radius, delta + radius]
    margin = relative_margin * abs(mean_b)
    return {"mean_new": mean_a, "mean_previous": mean_b, "difference": delta,
            "standard_error": se, "confidence": confidence, "confidence_interval": interval,
            "relative_difference": delta / mean_b if mean_b else None,
            "equivalence_margin": margin, "equivalent_within_margin": interval[0] > -margin and interval[1] < margin,
            "consistent_with_zero_difference": interval[0] <= 0 <= interval[1],
            "n_new": len(a), "n_previous": len(b), 'paired':paired,
            'degrees_of_freedom':int(df) if paired else float(df) if np.isfinite(df) else None}


def analysis_design(study):
    design=study['design']; sampling=study['sampling']
    return {'position_strata':design['stratum'],'stratum_weights':design['stratum_weights'],
            'replica_ids':list(range(sampling['replicas'])),'events_per_task':sampling['events_per_task'],
            'gun_energy_eV':study['fixed']['gun_energy_eV'],'electrode_weights':design['electrode_weights'],
            'design_hash':design['design_hash'],'injection_law':design.get('injection_law','uniform_surface')}


def score_output(result, prepared, task, *, time_resolved=False, score_hash=None):
    fixed=prepared['fixed']; totals=np.zeros(len(fixed['electrode_x_mm'])); n_hits=0; matrix=None
    surface=fixed['setSubThickness_um']*1e-6/2
    gap=fixed['setTopGap']
    for energy,x,y,z,arrival in records.read_hits(result['output']):
        count=physics.qp_counts(energy,x,y,z,arrival,gap,fixed['electrode_x_mm'],fixed['electrode_y_mm'],surface,time_resolved=time_resolved)
        if time_resolved:
            times,qps=count
            matrix=qps if matrix is None else matrix+qps
            totals+=qps.sum(axis=1)
        else: totals+=count
        n_hits+=len(energy)
    block={'position':task['position'],'replica':task['replica'],'events':task['events'],'total_qps':float(totals.sum()),
           'per_electrode_qps':totals.tolist(),'n_hits':n_hits}
    if time_resolved:
        if matrix is None:
            times=np.linspace(0,300000,1001); matrix=np.zeros((len(totals),1001))
        destination=Path(result['output']).with_name('qps.npz')
        np.savez_compressed(destination,time_ns=times,qps=matrix)
        block['time_data']=str(destination)
        block['time_data_sha256']=records.sha256_file(destination)
    block['score_hash']=score_hash or scoring_identity()
    return {**result,**block}


def density_response(blocks, prepared, defaults=None, cache=None, *, from_hits=False):
    """Stream pooled and realization time sums; retain the nonlinear pooled ODE."""
    ode=dict(defaults or {row[0]:row[1] for row in records.load_document(ROOT/'parameters.yaml')['morris']['QPDE_params']})
    ode.update(prepared.get('ode',{})); fixed=prepared['fixed']
    if not blocks: raise ValueError('density evolution needs complete simulation blocks')
    source,checksum=('output','sha256') if from_hits else ('time_data','time_data_sha256')
    geometry=(fixed['setTopGap'],tuple(fixed['electrode_x_mm']),tuple(fixed['electrode_y_mm']),fixed['setSubThickness_um']) if from_hits else None
    key=(from_hits,geometry,tuple((block[source],block[checksum],block['replica'],block['events']) for block in blocks))
    saved=cache.get(key) if cache is not None else None
    signatures=[]
    for block in blocks:
        stat=Path(block[source]).stat()
        signatures.append((stat.st_dev,stat.st_ino,stat.st_size,stat.st_mtime_ns,stat.st_ctime_ns))
    for i,block in enumerate(blocks):
        if saved is None or signatures[i]!=saved[3][i]:
            if records.sha256_file(block[source])!=block[checksum]: raise ValueError('raw hit checksum differs' if from_hits else 'time data checksum differs')
            stat=Path(block[source]).stat()
            if signatures[i]!=(stat.st_dev,stat.st_ino,stat.st_size,stat.st_mtime_ns,stat.st_ctime_ns):
                raise ValueError('density input changed while reading')
    if saved is None:
        replicas=sorted({block['replica'] for block in blocks}); sums={}; events={}; pooled=None; times=None
        for i,block in enumerate(blocks):
            if from_hits:
                time_grid=np.linspace(0,300000,1001); qps=np.zeros((len(fixed['electrode_x_mm']),1001))
                for columns in records.read_hits(block['output']):
                    time_grid,counts=physics.qp_counts(*columns,fixed['setTopGap'],fixed['electrode_x_mm'],fixed['electrode_y_mm'],fixed['setSubThickness_um']*1e-6/2,time_resolved=True)
                    qps+=counts
            else:
                with np.load(block['time_data'],allow_pickle=False) as data:
                    time_grid=data['time_ns']; qps=data['qps']
            stat=Path(block[source]).stat()
            if signatures[i]!=(stat.st_dev,stat.st_ino,stat.st_size,stat.st_mtime_ns,stat.st_ctime_ns):
                raise ValueError('density input changed while reading')
            if times is None: times=time_grid; pooled=np.zeros_like(qps)
            elif not np.array_equal(times,time_grid): raise ValueError('recorded time grids differ')
            replica=block['replica']
            if replica not in sums: sums[replica]=np.zeros_like(qps); events[replica]=0
            pooled+=qps; sums[replica]+=qps; events[replica]+=block['events']
        saved=(times,[pooled]+[sums[r] for r in replicas],[sum(events.values())]+[events[r] for r in replicas],signatures)
        if cache is not None:
            cache.clear(); cache[key]=saved
    elif saved[3]!=signatures:
        saved=(*saved[:3],signatures)
        cache[key]=saved
    times,arrays,event_counts=saved[:3]; outcomes=[]
    for qps,events in zip(arrays,event_counts):
        t,dg=physics.calculate_xqps(times,qps,fixed['setTopGap']/4.135667696e-15,ode['f_01'],ode['r'],ode['s'],ode['I_ph'],ode['pt'],ode['n_cooper'],fixed['electrode_height_um'],fixed['electrode_width_um'],fixed['setTopThickness_um'],events)
        outcomes.append(float(np.trapezoid(dg,t,axis=1).sum()))
    return {'integrated_density':outcomes[0],'integrated_density_by_replica':outcomes[1:],'density_parameters':ode,
            'density_sites':len({row['position'] for row in blocks})}


def analyse_results(study, rows):
    complete=[row for row in rows if row['status']=='complete']
    objective=study['analysis']['objective']; design=analysis_design(study); comparisons=[]
    reference=rows[0] if rows and rows[0]['status']=='complete' else None
    paired_requested=study['kind'] not in ('search','sensitivity') or (study['kind']=='search' and study['analysis'].get('diagnostics',False))
    if reference is not None and paired_requested:
        for candidate in complete:
            if candidate is reference: continue
            if objective=='integrated_density':
                delta=candidate['value']-reference['value']
                paired={'value':delta,'standard_error':None,'relative_to_b':delta/reference['value'] if reference['value'] else None,
                        'z':None,'units':candidate.get('units'),'uncertainty_status':'aggregate nonlinear density evolution has no measured paired uncertainty'}
            else:
                paired=paired_difference(candidate['blocks'],reference['blocks'],design,objective=objective,
                                                   bootstrap_draws=int(study['analysis'].get('bootstrap_draws',1000)))
            comparisons.append({'candidate':candidate['name'],'reference':reference['name'],**paired})
    geometry={key:value for key,value in study['fixed'].items() if key.startswith(('electrode_','setSub'))}
    report={'id':study['id'],'kind':study['kind'],'objective':objective,'events_total_per_candidate':study['events_total'],
            'sites':len(study['design']['sites_mm']),'replicas':study['sampling']['replicas'],
            'events_per_task':study['sampling']['events_per_task'],'design':design,'results':rows,'comparisons':comparisons,
            'geometry':geometry,'sites_mm':study['design']['sites_mm'],
            'streams':{k:study['sampling'][k] for k in ('seed_base','seed_bank','position_seed')},
            'complete':bool(rows) and len(complete)==len(rows),'limitations':['Calibrated interface model.','Property targets need measured-material validation.',
                                                            'QP yield is not a logical-error probability.']}
    if study['kind']=='sensitivity':
        recorded=study.get('morris_design') and all(all(key in row for key in ('inputs','trajectory','design_index')) for row in rows)
        if not recorded: raise ValueError('Morris reanalysis requires recorded sample inputs')
        candidates=sorted(rows,key=lambda row:row['design_index'])
        lookup={row['name']:row for row in rows}
        morris=study['morris_design']; problem=morris['problem']; names=problem['names']
        grouped={}
        for candidate in candidates: grouped.setdefault(candidate['trajectory'],[]).append(candidate)
        missing=[trajectory for trajectory,points in grouped.items()
                 if len(points)!=len(names)+1 or any(lookup.get(p['name'],{}).get('status')!='complete' for p in points)]
        if missing:
            report['sensitivity']={'status':'incomplete','missing_trajectories':missing,'declared_trajectories':len(grouped),
                                   'reason':'missing outcomes cannot be removed from Morris trajectories'}
        else:
            matrix=[[candidate['inputs'][name] for name in names] for candidate in candidates]
            response=[lookup[candidate['name']]['integrated_density'] if objective=='integrated_density' else lookup[candidate['name']]['value'] for candidate in candidates]
            method=study['analysis'].get('sensitivity_method','morris')
            result=sensitivity(matrix,response,names,problem=problem,method=method,
                                        log_outcome=study['analysis'].get('log_outcome',method=='pearson'),
                                        seed=morris['seed'],num_levels=morris['levels'])
            report['sensitivity']={**result,'status':'complete','response':objective,'morris_design':morris}
    if study['kind']=='tolerance':
        if reference is None:
            report['tolerance']={'locally_converged':False,'reason':'reference measurement is incomplete'}
        else:
            moves={}; paired={}; compared={row['candidate']:row for row in comparisons}
            for row in rows[1:]:
                name,side=row['name'].rsplit('_',1)
                label=name+(':minus' if side=='-1' else ':plus')
                moves[label]={'value':row.get('value')}
                if row['name'] in compared: paired[label]=compared[row['name']]
            for name in study.get('tolerance',{}):
                moves.setdefault(name+':minus',{'value':None}); moves.setdefault(name+':plus',{'value':None})
            varied={label.rsplit(':',1)[0] for label in moves}
            coordinates=physics.to_unit(reference['parameters'],study['space']) if varied<=set(study['space']['names']) else None
            report['tolerance']=tolerance_summary(reference,moves,paired=paired,coordinates=coordinates,names=study['space']['names'])
            report['tolerance']['untested_moves']=study.get('tolerance_skipped',[])
    return report


def read_analysis(path, output=None, settings=None):
    data=records.load_result(path,include_inputs=True); study=data['study']
    if data.get('raw_files_present') is False:
        raise ValueError('recover the recorded archive before raw-data reanalysis')
    expected=study.pop('resolved_hash')
    if records.content_hash(study)!=expected: raise ValueError('resolved study was changed after recording')
    study['resolved_hash']=expected
    if settings: study['analysis'].update(records.load_document(settings))
    declarations=data['candidate_inputs']; identity=scoring_identity(); design=analysis_design(study)
    defaults={entry[0]:entry[1] for entry in records.load_document(ROOT/'parameters.yaml')['morris']['QPDE_params']}
    cache={}
    compact_rows=[]
    paired_requested=study['kind'] not in ('search','sensitivity') or (study['kind']=='search' and study['analysis'].get('diagnostics',False))
    for row in data['results']:
        if row.get('status')!='complete': continue
        compact='blocks' not in row
        if compact: compact_rows.append(row)
        stored=records.load_result(path,include_inputs=True,candidate_id=row['candidate_id'])['results'][0] if compact else row
        blocks=[]; from_hits=False
        for block in stored['blocks']:
            if not Path(block['output']).is_file() or records.sha256_file(block['output'])!=block['sha256']:
                raise ValueError('saved measurements no longer match their raw files')
            if block.get('score_hash')!=identity:
                from_hits=True
                block={key:value for key,value in block.items() if key not in ('time_data','time_data_sha256')}
                block=score_output(block,declarations[row['candidate_id']],block,score_hash=identity)
            if not block.get('time_data') or not block.get('time_data_sha256'): from_hits=True
            blocks.append(block)
        counts=prepare_counts(blocks,design); options=study['analysis']; diagnostics=bool(options.get('diagnostics',study['kind']!='search'))
        if not compact or paired_requested: row['blocks']=blocks
        row['estimate']=summarize(counts,bootstrap_draws=int(options.get('bootstrap_draws',1000)),diagnostics=diagnostics)
        row['historical']=summarize(counts,objective='total_qps_per_primary',diagnostics=diagnostics)
        objective=options['objective']
        selected=row['estimate'] if objective in ('device_junction_qps_per_eV','device_weighted_junction_qps_per_energy') else row['historical']
        row.update(value=selected['value'],standard_error=selected['standard_error'],units=selected['units'])
        if objective=='integrated_density':
            row.update(density_response(blocks,declarations[row['candidate_id']],defaults,cache,from_hits=from_hits))
            row.update(value=row['integrated_density'],standard_error=None,units='MHz us')
    report=analyse_results(study,data['results'])
    for row in compact_rows: row.pop('blocks',None)
    if 'search_counts' in data:
        report['search_counts']=data['search_counts']
        if study.get('search',{}).get('measured_evaluations',False):
            report['complete']=bool(data['search_counts']['target_reached'])
            report['status']='complete' if report['complete'] else 'incomplete'
        if 'incomplete_reason' in data: report['incomplete_reason']=data['incomplete_reason']
    if output: records.write_json(output,report)
    return report


def compare_runs(first, second, output=None, margin=.05):
    a=records.load_result(first); b=records.load_result(second)
    if a['objective']!=b['objective'] or a['design']!=b['design']:
        raise ValueError('comparison requires the same objective and site/realization design')
    for key in ('geometry','sites_mm'):
        if key in a or key in b:
            if a.get(key)!=b.get(key): raise ValueError(f'comparison differs in physical {key}')
    if not np.isfinite(margin) or margin<=0: raise ValueError('relative equivalence margin must be positive and finite')
    rows=[]
    by_name={row['name']:row for row in b['results']}
    if set(by_name)!={row['name'] for row in a['results']}: raise ValueError('comparison requires matching candidate sets')
    for row in a['results']:
        other=by_name.get(row['name'])
        if row['status']!='complete' or not other or other['status']!='complete':
            raise ValueError('comparison requires complete matched candidates')
        if row['parameters']!=other['parameters']:
            raise ValueError('comparison material properties differ')
        if a['objective']=='integrated_density':
            rows.append({'name':row['name'],'difference':row['value']-other['value'],
                         'equivalent_within_margin':False,'uncertainty_status':'per-realization density evolution is required to measure uncertainty'})
            continue
        paired=paired_difference(row['blocks'],other['blocks'],a['design'],objective=a['objective'])
        ra=summarize(row['blocks'],a['design'],objective=a['objective'],diagnostics=False,uncertainty=False)['by_replica']
        rb=summarize(other['blocks'],b['design'],objective=b['objective'],diagnostics=False,uncertainty=False)['by_replica']
        matched=compare_repeated(ra,rb,relative_margin=margin,paired=True) if min(len(ra),len(rb))>=2 else {'status':'fewer than two realizations','equivalent_within_margin':False}
        repeated=compare_repeated(ra,rb,relative_margin=margin) if min(len(ra),len(rb))>=2 else {'status':'fewer than two realizations'}
        repeated['interpretation']='Welch interval ignores pairing; a diagnostic check, not a strict independent-sample experiment when random streams are shared.'
        rows.append({'name':row['name'],'paired':paired,'paired_realization_check':matched,'realization_check':repeated,
                     'equivalent_within_margin':matched['equivalent_within_margin']})
    report={'comparison':'selected response with matched sites and realizations','relative_equivalence_margin':margin,
            'objective':a['objective'],'same_random_streams':a.get('streams')==b.get('streams'),'results':rows}
    if output: records.write_json(output,report)
    return report


def plot_result(path, output, kind="summary", dpi=160, title=None):
    """Render PNG/PDF/SVG using only recorded data, never regenerated sites."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle

    data = records.load_result(path)
    kind = {"ranking": "materials", "material": "materials", "material-ranking": "materials",
            "orientation": "projection"}.get(kind, kind)
    allowed = ("summary", "layout", "hits", "time", "sensitivity", "convergence", "search", "projection", "materials")
    if kind not in allowed:
        raise ValueError(f"unknown figure kind: {kind}; choose {', '.join(allowed)}")
    fig, ax = plt.subplots(figsize=(8, 5.2), constrained_layout=True)
    try:
        if kind == "layout":
            geometry = data.get("geometry", data.get("study", {}).get("geometry", data.get("study", {}).get("fixed", data)))
            positions = geometry.get("electrodes_mm", geometry.get("electrode_xy_mm"))
            if positions is None:
                x = geometry.get("electrode_x_mm", geometry.get("electrode_x"))
                y = geometry.get("electrode_y_mm", geometry.get("electrode_y"))
                if x is None or y is None:
                    raise ValueError("layout needs recorded electrode coordinates in mm")
                positions = np.column_stack((x, y))
            electrodes = np.asarray(positions, dtype=float)
            ax.scatter(electrodes[:, 0], electrodes[:, 1], marker="s", s=70, label="electrodes")
            sites = data.get("sites_mm", data.get("design", {}).get("sites_mm", geometry.get("sites_mm", [])))
            if len(sites):
                sites = np.asarray(sites, dtype=float)
                ax.scatter(sites[:, 0], sites[:, 1], marker="*", s=40, label="injection sites")
            width = geometry.get("width_mm", geometry.get("chip_width_mm", geometry.get("setSubWidth_mm")))
            height = geometry.get("height_mm", geometry.get("chip_height_mm", geometry.get("setSubHeight_mm")))
            if width is not None and height is not None:
                ax.add_patch(Rectangle((-width / 2, -height / 2), width, height, fill=False))
            ax.set(xlabel="x [mm]", ylabel="y [mm]", aspect="equal")
            ax.legend(frameon=False)
        elif kind == "hits":
            counts = data.get("n_hits")
            if counts is None:
                counts = [row.get("result", {}).get("n_hits") for row in data.get("tasks", [])
                          if row.get("status") == "complete" and row.get("result") is not None]
                counts = [value for value in counts if value is not None]
                if not counts:
                    counts = [block["n_hits"] for row in data.get('results', []) if row.get("status", "complete") == "complete"
                              for block in row.get("blocks", []) if block.get("n_hits") is not None]
            if not len(counts) and data.get("hits_file"):
                counts = [sum(len(chunk[0]) for chunk in records.read_hits(data["hits_file"]))]
            if not len(counts):
                raise ValueError("hit figure needs measured completed-file hit counts")
            counts = np.asarray(counts, dtype=float)
            if (counts < 0).any() or not np.isfinite(counts).all():
                raise ValueError("invalid recorded hit counts")
            ax.hist(np.log1p(counts), bins=min(40, max(1, int(np.sqrt(len(counts))))))
            ticks = np.unique(np.rint(np.expm1(np.linspace(0, np.log1p(counts.max()), 6))).astype(int))
            ax.set_xticks(np.log1p(ticks), [f"{value:,}" for value in ticks])
            ax.set(xlabel="surface hits per complete file", ylabel="number of files")
            ax.text(.98, .98, f"{len(counts)} files; {np.count_nonzero(counts == 0)} measured zeros",
                    transform=ax.transAxes, ha="right", va="top")
        elif kind == "time":
            times = data.get("time_ns", data.get("snapshot_times_ns", data.get("times_ns")))
            values = data.get("qps_by_electrode_time", data.get("qps_time"))
            if times is None or values is None:
                sources = [block for row in data.get('results', []) if row.get("status", "complete") == "complete"
                           for block in row.get("blocks", []) if block.get("time_data")]
                if not sources:
                    sources = [row["result"] for row in data.get("tasks", [])
                               if row.get("status") == "complete" and row.get("result", {}).get("time_data")]
                for block in sources:
                    source = block["time_data"]
                    if block.get("time_data_sha256") and records.sha256_file(source) != block["time_data_sha256"]:
                        raise ValueError("recorded time data checksum differs")
                    recorded = records.load_result(source)
                    if times is None:
                        times = recorded["time_ns"]
                        values = recorded["qps_by_electrode_time"].copy()
                    elif np.array_equal(times, recorded["time_ns"]):
                        values += recorded["qps_by_electrode_time"]
                    else:
                        raise ValueError("recorded time grids differ")
            if times is None or values is None:
                raise ValueError("time figure needs recorded times_ns and qps_by_electrode_time")
            times = np.asarray(times) / 1000
            values = np.asarray(values)
            if values.ndim == 1:
                values = values[None, :]
            for index, series in enumerate(values):
                ax.plot(times, np.cumsum(series), label=f"electrode {index}", linewidth=1)
            ax.set(xlabel="arrival time [μs]", ylabel="cumulative quasiparticles")
            ax.legend(frameon=False, ncol=3, fontsize=7)
        elif kind == "sensitivity":
            result = data.get("sensitivity", data)
            names = result.get("names", result.get("parameters"))
            mu = result.get("mu_star")
            sigma = result.get("sigma")
            if names is None or mu is None or sigma is None:
                raise ValueError("sensitivity figure needs names, mu_star, and sigma")
            ax.scatter(mu, sigma)
            for x, y, label in zip(mu, sigma, names):
                ax.annotate(label, (x, y), xytext=(4, 4), textcoords="offset points", fontsize=8)
            ax.set(xlabel="mean absolute elementary effect, μ*", ylabel="standard deviation, σ")
        elif kind == "projection":
            rows = data.get("projection", data.get('results', []))
            if isinstance(rows, dict):
                rows = rows.get("materials", rows.get("rows", []))
            distances = [row["distance"] for row in rows if row.get("distance") is not None]
            labels = [str(row.get('id',row.get('name',row.get('material',index))))+(' ['+str(row['matched'])+'/'+str(row['total'])+']' if 'matched' in row else '') for index,row in enumerate(rows)
                      if row.get("distance") is not None]
            if not distances:
                raise ValueError("projection figure needs measured-material distances")
            indices = np.arange(len(distances))
            ax.barh(np.arange(len(indices)), np.asarray(distances)[indices])
            ax.set_yticks(np.arange(len(indices)), np.asarray(labels)[indices])
            ax.invert_yaxis()
            ax.set(xlabel=data.get("distance_label", "distance in recorded property coordinates"))
        else:
            rows = [row for row in data.get('results', []) if row.get("kind", "measured") in ("measured", "physical_zero")
                    and row.get("status", "complete") in ("complete", "done", "ok")
                    and not any(row.get(key)=="invalid" for key in ("validity","validity_state","registered_validity"))]
            if not rows and any(name in data for name in ("mean", "estimate", "yield")):
                rows = [data]
            if not rows:
                raise ValueError("figure needs complete measured observations")
            recorded_units = {row["units"] for row in rows if row.get("units")}
            if len(recorded_units) > 1:
                raise ValueError("cannot compare responses with different recorded units")
            units = next(iter(recorded_units), data.get("units"))
            objective = data.get("objective", "")
            if units == "summed_junction_QPs_per_primary" or objective in ("total_qps_per_primary", "total_QPs_per_primary", "junction_qps_per_primary"):
                label = "summed junction QPs / primary"
            elif units == "weighted_junction_QPs_per_injected_eV" or objective in ("device_junction_qps_per_eV", "device_weighted_junction_qps_per_energy"):
                label = "weighted junction QPs / injected eV"
            elif units == "MHz us" or objective == "integrated_density":
                label = "integrated decoherence rate [MHz μs]"
            else:
                label = units.replace("_", " ") if units else "recorded response"
            label = data.get("yield_label", label)
            means = np.asarray([row['value'] for row in rows])
            errors=np.asarray([row['standard_error'] if row.get('standard_error') is not None else np.nan for row in rows])
            if not np.isfinite(means).all() or np.isinf(errors).any() or (errors < 0).any():
                raise ValueError("invalid recorded means or uncertainties")
            if np.isnan(errors).any(): ax.text(.02,.02,"uncertainty unavailable for unmarked errors",transform=ax.transAxes,fontsize=8)
            if kind == "convergence":
                x = [row.get("n_sites", row.get("events", row.get("n", index + 1))) for index, row in enumerate(rows)]
                ax.errorbar(x, means, yerr=errors, marker="o", capsize=3)
                ax.set_xlabel(data.get("x_label", "sample count"))
            elif kind == "search":
                groups = sorted(set(str(row.get("method", "search")) for row in rows))
                for group in groups:
                    chosen = [index for index, row in enumerate(rows) if str(row.get("method", "search")) == group]
                    x = [rows[index].get("evaluation", order + 1) for order, index in enumerate(chosen)]
                    ax.plot(x, np.minimum.accumulate(means[chosen]), marker=".", label=group)
                ax.set_xlabel("complete candidate evaluations")
                ax.legend(frameon=False)
            else:
                indices = np.argsort(means)
                labels = [str(row.get("name", row.get("material", row.get("candidate_id", index)))) for index, row in enumerate(rows)]
                ax.errorbar(means[indices], np.arange(len(indices)), xerr=errors[indices], fmt="o", capsize=3)
                ax.set_yticks(np.arange(len(indices)), np.asarray(labels)[indices], fontsize=8)
                ax.invert_yaxis()
                ax.set_xlabel(label)
            if kind in ("convergence", "search"):
                ax.set_ylabel(label)
        ax.set_title(title or data.get("title", kind.replace("_", " ").capitalize()))
        ax.grid(alpha=.2)
        target = Path(output)
        if target.suffix.lower() not in (".png", ".pdf", ".svg"):
            raise ValueError("figure output must be PNG, PDF, or SVG")
        target.parent.mkdir(parents=True, exist_ok=True)
        fig.savefig(target, dpi=dpi)
    finally:
        plt.close(fig)
    return target


def export_presentation(figures, output, title="Material-property study"):
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as error:
        raise RuntimeError("presentation export needs python-pptx") from error
    presentation = Presentation()
    for index, path in enumerate(figures):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title if len(figures) == 1 else f"{title}: {index + 1}"
        slide.shapes.add_picture(str(path), Inches(.3), Inches(1), width=Inches(9.3))
    target = Path(output)
    target.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(target)
    return target


def summarize(blocks, design=None, *, objective='device_junction_qps_per_eV', diagnostics=True,
              uncertainty=True, bootstrap_draws=1000, bootstrap_seed=7,
              tail_bootstrap_draws=500, tail_bootstrap_seed=11, robustness=1., alpha=1., beta=1.):
    """One selected response and its error from Q[site, replica, electrode]."""
    data=prepare_counts(blocks,design); design=data['design']; totals=data['totals'].T
    events=design['events_per_task']; historical=objective in ('total_qps','total_qps_per_primary','total_QPs_per_primary',
        'junction_qps_per_primary','robust_mean_plus_se','max_electrode_qps_per_primary','weighted_sum')
    if historical:
        n_sim=events*totals.size; q=totals.ravel()/events
        se_total=float(np.sqrt(totals.shape[0]*totals.var(axis=0,ddof=1).sum())) if totals.shape[0]>1 else None
        se_block=float(q.std(ddof=1)/np.sqrt(q.size)) if q.size>1 else 0.
        value=float(totals.mean()/events)
        output={'total_qps':float(totals.sum()),'total_qps_per_primary':value,'junction_qps_per_primary':value,
                'total_qps_standard_error':se_total,'total_qps_per_primary_standard_error':se_total/n_sim if se_total is not None else None,
                'robust_mean_plus_se':float(q.mean()+robustness*se_block),'across_block_standard_error':se_block,
                'n_sites':totals.shape[1],'n_replicas':totals.shape[0],'by_replica':(totals.mean(axis=1)/events).tolist()}
        if diagnostics:
            rng=np.random.default_rng(12345); site_totals=totals.sum(axis=0)
            draws=site_totals[rng.integers(0,len(site_totals),size=(2000,len(site_totals)))].sum(axis=1)
            output.update(site_bootstrap_ci_total=np.percentile(draws,[2.5,97.5]).tolist(),
                          p90_position=float(np.percentile(totals.mean(axis=0)/events,90)))
        if data['counts'] is not None:
            electrode_totals=data['counts'].sum(axis=(0,1))
            output['max_electrode_qps_per_primary']=float(electrode_totals.max()/n_sim)
            output['weighted_sum']=alpha*value+beta*output['max_electrode_qps_per_primary']
        key='total_qps_per_primary' if objective=='total_QPs_per_primary' else objective
        if key not in output: raise ValueError(f'unknown or unavailable historical objective {objective!r}')
        se=se_total if key=='total_qps' else se_block if key=='robust_mean_plus_se' else se_total/n_sim if key in ('total_qps_per_primary','junction_qps_per_primary') and se_total is not None else None
        output.update(value=output[key],standard_error=se if uncertainty else None,objective=objective,
                      units='summed_junction_QPs_per_primary' if key!='total_qps' else 'summed_junction_QPs')
        return output
    if objective not in ('device_junction_qps_per_eV','device_weighted_junction_qps_per_energy') or data['weighted'] is None:
        raise ValueError(f'per-electrode response is unavailable for objective {objective!r}')
    burden=data['weighted']; regions=[(label,positions,weight,burden[positions]) for label,positions,weight in data['groups']]
    per_site=burden.mean(axis=1)
    means={label:float(per_site[positions].mean()) for label,positions,_,_ in regions}
    value=sum(weight*means[label] for label,_,weight,_ in regions)
    se=_resample(regions,bootstrap_draws,bootstrap_seed) if uncertainty else None
    output={'value':float(value),'standard_error':se,'relative_standard_error':se/value if se is not None and value else None,
            'means_by_stratum':means,'sites_by_stratum':{label:len(positions) for label,positions,_,_ in regions},
            'weights_by_stratum':dict(design['stratum_weights']),'naive_equal_site_mean':float(per_site.mean()),
            'maximum_site_leverage':float(max(weight*per_site[i]/len(positions)/value if value else 0. for _,positions,weight,_ in regions for i in positions)),
            'n_sites':len(per_site),'n_replicas':burden.shape[1],'events_per_task':events,'design_hash':design['design_hash'],
            'by_replica':np.sum([weight*matrix.mean(axis=0) for _,_,weight,matrix in regions],axis=0).tolist(),
            'units':'weighted_junction_QPs_per_injected_eV'}
    if diagnostics:
        worst=data['worst']; per_site=worst.mean(axis=1)
        output['spatial_r95']=_tail_mean([(per_site[i],weight/len(positions)) for _,positions,weight,_ in regions for i in positions])
        output['spatial_r95_standard_error']=_resample([(label,positions,weight,worst[positions]) for label,positions,weight,_ in regions],tail_bootstrap_draws,tail_bootstrap_seed,tail=True) if uncertainty else None
    return output


def scoring_identity():
    return records.content_hash({'physics':records.sha256_file(physics.__file__),
                                 'calculation':inspect.getsource(score_output)})
