"""Figures from measured simulation records."""

import argparse
import csv
import json
import sys
from pathlib import Path

import numpy as np

if not __package__:
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import records


def plot_result(path, output, kind="summary", dpi=160, title=None):
    """Render PNG/PDF/SVG using only recorded data, never regenerated sites."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle

    data = records.load_result(path)
    kind = {"ranking": "materials", "material": "materials", "material-ranking": "materials",
            "orientation": "projection"}.get(kind, kind)
    allowed = ("summary", "layout", "hits", "time", "sensitivity", "convergence", "search", "projection", "materials")
    if kind not in allowed:
        raise ValueError(f"unknown figure kind: {kind}; choose {', '.join(allowed)}")
    fig, ax = plt.subplots(figsize=(8, 5.2), constrained_layout=True)
    try:
        if kind == "layout":
            geometry = data.get("geometry", data.get("study", {}).get("geometry", data.get("study", {}).get("fixed", data)))
            positions = geometry.get("electrodes_mm", geometry.get("electrode_xy_mm"))
            if positions is None:
                x = geometry.get("electrode_x_mm", geometry.get("electrode_x"))
                y = geometry.get("electrode_y_mm", geometry.get("electrode_y"))
                if x is None or y is None:
                    raise ValueError("layout needs recorded electrode coordinates in mm")
                positions = np.column_stack((x, y))
            electrodes = np.asarray(positions, dtype=float)
            ax.scatter(electrodes[:, 0], electrodes[:, 1], marker="s", s=70, label="electrodes")
            sites = data.get("sites_mm", data.get("design", {}).get("sites_mm", geometry.get("sites_mm", [])))
            if len(sites):
                sites = np.asarray(sites, dtype=float)
                ax.scatter(sites[:, 0], sites[:, 1], marker="*", s=40, label="injection sites")
            width = geometry.get("width_mm", geometry.get("chip_width_mm", geometry.get("setSubWidth_mm")))
            height = geometry.get("height_mm", geometry.get("chip_height_mm", geometry.get("setSubHeight_mm")))
            if width is not None and height is not None:
                ax.add_patch(Rectangle((-width / 2, -height / 2), width, height, fill=False))
            ax.set(xlabel="x [mm]", ylabel="y [mm]", aspect="equal")
            ax.legend(frameon=False)
        elif kind == "hits":
            counts = data.get("n_hits")
            if counts is None:
                counts = [row.get("result", {}).get("n_hits") for row in data.get("tasks", [])
                          if row.get("status") == "complete" and row.get("result") is not None]
                counts = [value for value in counts if value is not None]
                if not counts:
                    counts = [block["n_hits"] for row in data.get('results', []) if row.get("status", "complete") == "complete"
                              for block in row.get("blocks", []) if block.get("n_hits") is not None]
            if not len(counts) and data.get("hits_file"):
                counts = [sum(len(chunk[0]) for chunk in records.read_hits(data["hits_file"]))]
            if not len(counts):
                raise ValueError("hit figure needs measured completed-file hit counts")
            counts = np.asarray(counts, dtype=float)
            if (counts < 0).any() or not np.isfinite(counts).all():
                raise ValueError("invalid recorded hit counts")
            if np.any(counts != np.floor(counts)):
                raise ValueError("surface-hit records must be integer counts")
            from matplotlib.ticker import MaxNLocator
            logarithmic=data.get('histogram_log_x',True)
            bins=data.get('histogram_bins',min(40,max(1,int(np.sqrt(len(counts))))))
            if isinstance(bins,str) and bins=='auto':
                bins=np.arange(counts.min()-.5,counts.max()+1.5) if not logarithmic and counts.max()-counts.min()<=100 else min(40,max(1,int(np.sqrt(len(counts)))))
            if isinstance(bins,(int,np.integer)) and bins<1: raise ValueError('histogram bins must be positive')
            ax.hist(np.log1p(counts) if logarithmic else counts,bins=bins,color='#30658a',edgecolor='white',linewidth=.8)
            if logarithmic:
                ticks = np.unique(np.rint(np.expm1(np.linspace(0, np.log1p(counts.max()), 6))).astype(int))
                ax.set_xticks(np.log1p(ticks), [f"{value:,}" for value in ticks])
            else: ax.xaxis.set_major_locator(MaxNLocator(integer=True))
            ax.yaxis.set_major_locator(MaxNLocator(integer=True))
            ax.set(xlabel="surface-hit records per file", ylabel="number of files")
            ax.text(.98, .98, f"{len(counts)} files; {int(counts.sum()):,} records\n{np.count_nonzero(counts == 0)} zero-hit files",
                    transform=ax.transAxes, ha="right", va="top")
        elif kind == "time":
            times = data.get("time_ns", data.get("snapshot_times_ns", data.get("times_ns")))
            values = data.get("qps_by_electrode_time", data.get("qps_time"))
            if times is None or values is None:
                sources = [block for row in data.get('results', []) if row.get("status", "complete") == "complete"
                           for block in row.get("blocks", []) if block.get("time_data")]
                if not sources:
                    sources = [row["result"] for row in data.get("tasks", [])
                               if row.get("status") == "complete" and row.get("result", {}).get("time_data")]
                for block in sources:
                    source = block["time_data"]
                    if block.get("time_data_sha256") and records.sha256_file(source) != block["time_data_sha256"]:
                        raise ValueError("recorded time data checksum differs")
                    recorded = records.load_result(source)
                    if times is None:
                        times = recorded["time_ns"]
                        values = recorded["qps_by_electrode_time"].copy()
                    elif np.array_equal(times, recorded["time_ns"]):
                        values += recorded["qps_by_electrode_time"]
                    else:
                        raise ValueError("recorded time grids differ")
            if times is None or values is None:
                raise ValueError("time figure needs recorded times_ns and qps_by_electrode_time")
            times = np.asarray(times) / 1000
            values = np.asarray(values)
            if values.ndim == 1:
                values = values[None, :]
            for index, series in enumerate(values):
                ax.plot(times, np.cumsum(series), label=f"electrode {index}", linewidth=1)
            ax.set(xlabel="arrival time [μs]", ylabel="cumulative quasiparticles")
            ax.legend(frameon=False, ncol=3, fontsize=7)
        elif kind == "sensitivity":
            result = data.get("sensitivity", data)
            names = result.get("names", result.get("parameters"))
            mu = result.get("mu_star")
            sigma = result.get("sigma")
            if names is None or mu is None or sigma is None:
                raise ValueError("sensitivity figure needs names, mu_star, and sigma")
            ax.scatter(mu, sigma)
            for x, y, label in zip(mu, sigma, names):
                ax.annotate(label, (x, y), xytext=(4, 4), textcoords="offset points", fontsize=8)
            ax.set(xlabel="mean absolute elementary effect, μ*", ylabel="standard deviation, σ")
        elif kind == "projection":
            rows = data.get("projection", data.get('results', []))
            if isinstance(rows, dict):
                rows = rows.get("materials", rows.get("rows", []))
            distances = [row["distance"] for row in rows if row.get("distance") is not None]
            labels = [str(row.get('id',row.get('name',row.get('material',index))))+(' ['+str(row['matched'])+'/'+str(row['total'])+']' if 'matched' in row else '') for index,row in enumerate(rows)
                      if row.get("distance") is not None]
            if not distances:
                raise ValueError("projection figure needs measured-material distances")
            indices = np.arange(len(distances))
            ax.barh(np.arange(len(indices)), np.asarray(distances)[indices])
            ax.set_yticks(np.arange(len(indices)), np.asarray(labels)[indices])
            ax.invert_yaxis()
            ax.set(xlabel=data.get("distance_label", "distance in recorded property coordinates"))
        else:
            rows = [row for row in data.get('results', []) if row.get("kind", "measured") in ("measured", "physical_zero")
                    and row.get("status", "complete") in ("complete", "done", "ok")
                    and not any(row.get(key)=="invalid" for key in ("validity","validity_state","registered_validity"))]
            if not rows and any(name in data for name in ("mean", "estimate", "yield")):
                rows = [data]
            if not rows:
                raise ValueError("figure needs complete measured observations")
            recorded_units = {row["units"] for row in rows if row.get("units")}
            if len(recorded_units) > 1:
                raise ValueError("cannot compare responses with different recorded units")
            units = next(iter(recorded_units), data.get("units"))
            objective = data.get("objective", "")
            if units == "summed_junction_QPs_per_primary" or objective in ("total_qps_per_primary", "total_QPs_per_primary", "junction_qps_per_primary"):
                label = "summed junction QPs / primary"
            elif units == "weighted_junction_QPs_per_injected_eV" or objective in ("device_junction_qps_per_eV", "device_weighted_junction_qps_per_energy"):
                label = "weighted junction QPs / injected eV"
            elif units == "MHz us" or objective == "integrated_density":
                label = "integrated decoherence rate [MHz μs]"
            else:
                label = units.replace("_", " ") if units else "recorded response"
            label = data.get("yield_label", label)
            means = np.asarray([row['value'] for row in rows])
            errors=np.asarray([row['standard_error'] if row.get('standard_error') is not None else np.nan for row in rows])
            if not np.isfinite(means).all() or np.isinf(errors).any() or (errors < 0).any():
                raise ValueError("invalid recorded means or uncertainties")
            if np.isnan(errors).any(): ax.text(.02,.02,"uncertainty unavailable for unmarked errors",transform=ax.transAxes,fontsize=8)
            if kind == "convergence":
                x = [row.get("n_sites", row.get("events", row.get("n", index + 1))) for index, row in enumerate(rows)]
                ax.errorbar(x, means, yerr=errors, marker="o", capsize=3)
                ax.set_xlabel(data.get("x_label", "sample count"))
            elif kind == "search":
                groups = sorted(set(str(row.get("method", "search")) for row in rows))
                for group in groups:
                    chosen = [index for index, row in enumerate(rows) if str(row.get("method", "search")) == group]
                    x = [rows[index].get("evaluation", order + 1) for order, index in enumerate(chosen)]
                    ax.plot(x, np.minimum.accumulate(means[chosen]), marker=".", label=group)
                ax.set_xlabel("complete candidate evaluations")
                ax.legend(frameon=False)
            else:
                indices = np.argsort(means)
                labels = [str(row.get("name", row.get("material", row.get("candidate_id", index)))) for index, row in enumerate(rows)]
                ax.errorbar(means[indices], np.arange(len(indices)), xerr=errors[indices], fmt="o", capsize=3)
                ax.set_yticks(np.arange(len(indices)), np.asarray(labels)[indices], fontsize=8)
                ax.invert_yaxis()
                ax.set_xlabel(label)
            if kind in ("convergence", "search"):
                ax.set_ylabel(label)
        ax.set_title(title or data.get("title", kind.replace("_", " ").capitalize()))
        ax.grid(alpha=.2)
        target = Path(output)
        if target.suffix.lower() not in (".png", ".pdf", ".svg"):
            raise ValueError("figure output must be PNG, PDF, or SVG")
        target.parent.mkdir(parents=True, exist_ok=True)
        fig.savefig(target, dpi=dpi)
    finally:
        plt.close(fig)
    return target


def export_presentation(figures, output, title="Material-property study"):
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as error:
        raise RuntimeError("presentation export needs python-pptx") from error
    presentation = Presentation()
    for index, path in enumerate(figures):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title if len(figures) == 1 else f"{title}: {index + 1}"
        slide.shapes.add_picture(str(path), Inches(.3), Inches(1), width=Inches(9.3))
    target = Path(output)
    target.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(target)
    return target


def count_hit_files(paths, *, expected_checksums=None):
    """H_f=sum_k 1 over CSV data rows in file f; both surfaces, including zero E_dep.

    Header-only files have H_f=0. A malformed or changing file has no count.
    """
    files=[]; seen=set()
    for path in paths:
        path=Path(path).expanduser().resolve(strict=True)
        if path in seen: continue
        seen.add(path); before=path.stat()
        count=sum(len(columns[0]) for columns in records.read_hits(path))
        if expected_checksums and path in expected_checksums:
            if records.sha256_file(path)!=expected_checksums[path]: raise ValueError(f'hit-file checksum differs: {path}')
        after=path.stat()
        if (before.st_dev,before.st_ino,before.st_size,before.st_mtime_ns,before.st_ctime_ns)!=(after.st_dev,after.st_ino,after.st_size,after.st_mtime_ns,after.st_ctime_ns):
            raise RuntimeError(f'hit file changed while counting: {path}')
        files.append({'path':str(path),'surface_hit_records':count})
    if not files: raise ValueError('no complete hit files found')
    counts=np.asarray([row['surface_hit_records'] for row in files],dtype=np.int64)
    return {'files':files,'n_files':len(files),'surface_hit_records':int(counts.sum()),
            'zero_hit_files':int(np.count_nonzero(counts==0)),'min_records':int(counts.min()),
            'max_records':int(counts.max()),'mean_records':float(counts.mean()),'median_records':float(np.median(counts))}


def plot_hit_histogram(report, output, *, bins='auto', log_x=False, title=None):
    """Histogram of H_f against file counts; integer bins keep measured zeros."""
    counts=[row['surface_hit_records'] for row in report['files']]
    if not counts or len(counts)!=report['n_files'] or sum(counts)!=report['surface_hit_records']:
        raise ValueError('per-file counts differ from the stated totals')
    data={'n_hits':counts,'histogram_bins':bins,'histogram_log_x':log_x}
    return plot_result(data,output,'hits',title=title or 'Surface-hit records in completed files')


def main(argv=None):
    parser=argparse.ArgumentParser(description='Count surface-hit records and plot saved measurements.')
    commands=parser.add_subparsers(dest='command',required=True)
    hits=commands.add_parser('hits',help='Count rows in complete hit files; plot records per file against number of files.')
    hits.add_argument('inputs',nargs='+',type=Path,help='Hit files, raw-hit directories, or recorded study directories.')
    hits.add_argument('--output',type=Path,required=True)
    hits.add_argument('--counts',type=Path,help='Write one count per file as CSV.')
    hits.add_argument('--pattern',action='append',help='Raw-directory glob; default: hits.csv and *_hitsfile.txt.')
    hits.add_argument('--bins',type=int)
    hits.add_argument('--log-x',action='store_true',help='Space bins in log(1+H_f); retain zero-hit files.')
    hits.add_argument('--title')
    result=commands.add_parser('result',help='Plot a recorded response, sensitivity, geometry, or search.')
    result.add_argument('path');result.add_argument('--output',required=True);result.add_argument('--kind',default='summary')
    args=parser.parse_args(argv)
    try:
        if args.command=='result':
            print(json.dumps({'figure':str(plot_result(args.path,args.output,args.kind))}));return 0
        paths=[];checksums={}
        for source in args.inputs:
            source=source.expanduser().resolve(strict=True)
            if source.is_file():
                if 'tasks' in source.parts and not (source.parent/'complete').is_file():
                    raise ValueError(f'unfinished transport file: {source}')
                paths.append(source);continue
            studies=[source] if (source/'results.sqlite').is_file() else [p.parent for p in (source/'studies').glob('*/results.sqlite')]
            if studies:
                for study in studies:
                    with records.Results(study/'results.sqlite',create=False) as database:
                        origin=database.read('state','output_root')
                        for row in database.connection.execute("SELECT result_json FROM tasks WHERE status='complete'"):
                            measured=json.loads(row[0]);path=Path(measured['output'])
                            if origin and path.is_relative_to(origin): path=study/path.relative_to(origin)
                            path=path.resolve(strict=True)
                            if not measured.get('sha256'): raise ValueError(f'missing recorded hit-file checksum: {path}')
                            paths.append(path);checksums[path]=measured['sha256']
            else:
                for pattern in args.pattern or ('hits.csv','*_hitsfile.txt'):
                    paths.extend(p for p in source.rglob(pattern) if p.is_file() and ('tasks' not in p.parts or (p.parent/'complete').is_file()))
        selected={path.resolve() for path in paths}
        if args.output.resolve() in selected or (args.counts and args.counts.resolve() in selected):
            raise ValueError('figure and count outputs must differ from the input hit files')
        if args.counts and args.output.resolve()==args.counts.resolve():
            raise ValueError('figure and count table require different output paths')
        report=count_hit_files(sorted(paths),expected_checksums=checksums)
        figure=plot_hit_histogram(report,args.output,bins=args.bins if args.bins is not None else 'auto',log_x=args.log_x,title=args.title)
        if args.counts:
            args.counts.parent.mkdir(parents=True,exist_ok=True)
            with args.counts.open('w',newline='') as handle:
                writer=csv.DictWriter(handle,fieldnames=('path','surface_hit_records'),lineterminator='\n');writer.writeheader();writer.writerows(report['files'])
        print(json.dumps({**{key:value for key,value in report.items() if key!='files'},'figure':str(figure),
                          'counts_csv':str(args.counts) if args.counts else None},indent=2))
        return 0
    except (ValueError,OSError,RuntimeError) as error:
        parser.error(str(error))


if __name__=='__main__':
    raise SystemExit(main())
