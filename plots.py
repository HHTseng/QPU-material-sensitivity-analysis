"""Figures from recorded geometry, hits, and scientific results."""

import json
from pathlib import Path

import numpy as np

from records import Results, read_hits, sha256_file


def load_result(path):
    if isinstance(path, dict):
        return path
    source = Path(path)
    if source.is_dir():
        choices = [source / name for name in ("summary.json", "analysis.json", "results.json", "results.sqlite", "experiment.sqlite")]
        source = next((item for item in choices if item.is_file()), None)
        if source is None:
            raise ValueError(f"no recorded result in {path}")
    if source.suffix in (".sqlite", ".sqlite3", ".db"):
        with Results(source, create=False) as database:
            if "observations" in database.tables:
                rows = database.observations()
                data = {"rows": [{"candidate_id": row["candidate_id"],
                                  "kind": row["kind"], **(row["value"] if isinstance(row["value"], dict)
                                                         else {"mean": row["value"]})} for row in rows],
                        "tasks": database.tasks()}
                experiments = database.rows("experiments")
                if experiments:
                    data["study"] = json.loads(experiments[0]["input_json"])
                return _bind_files(data, source.parent, database.state("output_root"))
            trials = database.rows("trials") if "trials" in database.tables else []
            return {"rows": trials, "source": str(source)}
    if source.suffix == ".csv":
        return {"hits_file": str(source)}
    if source.suffix == ".npz":
        with np.load(source, allow_pickle=False) as values:
            return {"time_ns": values["time_ns"], "qps_by_electrode_time": values["qps"]}
    with source.open(encoding="utf-8") as handle:
        data = json.load(handle)
    data = data if isinstance(data, dict) else {"rows": data}
    if (source.parent / "results.sqlite").is_file():
        with Results(source.parent / "results.sqlite", create=False) as database:
            origin = database.state("output_root") if "state" in database.tables else None
        return _bind_files(data, source.parent, origin)
    return data


def _bind_files(data, root, origin=None):
    blocks = [block for row in _rows(data) for block in row.get("blocks", [])]
    blocks += [row["result"] for row in data.get("tasks", []) if row.get("result") is not None]
    for block in blocks:
        recorded = block.get("output", "")
        old_root = Path(origin) if origin else Path(recorded.split("/tasks/", 1)[0])
        if not origin and "/tasks/" not in recorded:
            continue
        for field in ("output", "time_data"):
            if field in block:
                relative = Path(block[field]).relative_to(old_root)
                if ".." in relative.parts:
                    raise ValueError("recorded result path escapes its output directory")
                block[field] = str(Path(root) / relative)
    return data


def _rows(data):
    rows = data.get("rows", data.get("results", data.get("observations", data.get("candidates", []))))
    if isinstance(rows, dict):
        rows = [{"name": name, **(value if isinstance(value, dict) else {"mean": value})}
                for name, value in rows.items()]
    return [{**row, **row.get("value", {})} if isinstance(row.get("value"), dict) else row for row in rows]


def _estimate(row):
    for name in ("mean", "estimate", "yield", "objective", "J", "qp_mean", "mean_qp_per_eV", "value"):
        value = row.get(name)
        if value is not None and np.isscalar(value):
            return float(value)
    raise ValueError("record is missing a measured mean")


def _uncertainty(row):
    for name in ("standard_error", "se", "stderr", "error"):
        if row.get(name) is not None:
            return float(row[name])
    return 0.0


def plot_result(path, output, kind="summary", dpi=160, title=None):
    """Render PNG/PDF/SVG using only recorded data, never regenerated sites."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle

    data = load_result(path)
    kind = {"ranking": "materials", "material": "materials", "material-ranking": "materials",
            "orientation": "projection"}.get(kind, kind)
    allowed = ("summary", "layout", "hits", "time", "sensitivity", "convergence", "search", "projection", "materials")
    if kind not in allowed:
        raise ValueError(f"unknown figure kind: {kind}; choose {', '.join(allowed)}")
    fig, ax = plt.subplots(figsize=(8, 5.2), constrained_layout=True)
    try:
        if kind == "layout":
            geometry = data.get("geometry", data.get("study", {}).get("geometry", data.get("study", {}).get("fixed", data)))
            positions = geometry.get("electrodes_mm", geometry.get("electrode_xy_mm"))
            if positions is None:
                x = geometry.get("electrode_x_mm", geometry.get("electrode_x"))
                y = geometry.get("electrode_y_mm", geometry.get("electrode_y"))
                if x is None or y is None:
                    raise ValueError("layout needs recorded electrode coordinates in mm")
                positions = np.column_stack((x, y))
            electrodes = np.asarray(positions, dtype=float)
            ax.scatter(electrodes[:, 0], electrodes[:, 1], marker="s", s=70, label="electrodes")
            sites = data.get("sites_mm", data.get("design", {}).get("sites_mm", geometry.get("sites_mm", [])))
            if len(sites):
                sites = np.asarray(sites, dtype=float)
                ax.scatter(sites[:, 0], sites[:, 1], marker="*", s=40, label="injection sites")
            width = geometry.get("width_mm", geometry.get("chip_width_mm", geometry.get("setSubWidth_mm")))
            height = geometry.get("height_mm", geometry.get("chip_height_mm", geometry.get("setSubHeight_mm")))
            if width is not None and height is not None:
                ax.add_patch(Rectangle((-width / 2, -height / 2), width, height, fill=False))
            ax.set(xlabel="x [mm]", ylabel="y [mm]", aspect="equal")
            ax.legend(frameon=False)
        elif kind == "hits":
            counts = data.get("n_hits")
            if counts is None:
                counts = [row.get("result", {}).get("n_hits") for row in data.get("tasks", [])
                          if row.get("status") == "complete" and row.get("result") is not None]
                counts = [value for value in counts if value is not None]
                if not counts:
                    counts = [block["n_hits"] for row in _rows(data) if row.get("status", "complete") == "complete"
                              for block in row.get("blocks", []) if block.get("n_hits") is not None]
            if not len(counts) and data.get("hits_file"):
                counts = [sum(len(chunk[0]) for chunk in read_hits(data["hits_file"]))]
            if not len(counts):
                raise ValueError("hit figure needs measured completed-file hit counts")
            counts = np.asarray(counts, dtype=float)
            if (counts < 0).any() or not np.isfinite(counts).all():
                raise ValueError("invalid recorded hit counts")
            ax.hist(np.log1p(counts), bins=min(40, max(1, int(np.sqrt(len(counts))))))
            ticks = np.unique(np.rint(np.expm1(np.linspace(0, np.log1p(counts.max()), 6))).astype(int))
            ax.set_xticks(np.log1p(ticks), [f"{value:,}" for value in ticks])
            ax.set(xlabel="surface hits per complete file", ylabel="number of files")
            ax.text(.98, .98, f"{len(counts)} files; {np.count_nonzero(counts == 0)} measured zeros",
                    transform=ax.transAxes, ha="right", va="top")
        elif kind == "time":
            times = data.get("time_ns", data.get("snapshot_times_ns", data.get("times_ns")))
            values = data.get("qps_by_electrode_time", data.get("qps_time"))
            if times is None or values is None:
                sources = [block for row in _rows(data) if row.get("status", "complete") == "complete"
                           for block in row.get("blocks", []) if block.get("time_data")]
                if not sources:
                    sources = [row["result"] for row in data.get("tasks", [])
                               if row.get("status") == "complete" and row.get("result", {}).get("time_data")]
                for block in sources:
                    source = block["time_data"]
                    if block.get("time_data_sha256") and sha256_file(source) != block["time_data_sha256"]:
                        raise ValueError("recorded time data checksum differs")
                    recorded = load_result(source)
                    if times is None:
                        times = recorded["time_ns"]
                        values = recorded["qps_by_electrode_time"].copy()
                    elif np.array_equal(times, recorded["time_ns"]):
                        values += recorded["qps_by_electrode_time"]
                    else:
                        raise ValueError("recorded time grids differ")
            if times is None or values is None:
                raise ValueError("time figure needs recorded times_ns and qps_by_electrode_time")
            times = np.asarray(times) / 1000
            values = np.asarray(values)
            if values.ndim == 1:
                values = values[None, :]
            for index, series in enumerate(values):
                ax.plot(times, np.cumsum(series), label=f"electrode {index}", linewidth=1)
            ax.set(xlabel="arrival time [μs]", ylabel="cumulative quasiparticles")
            ax.legend(frameon=False, ncol=3, fontsize=7)
        elif kind == "sensitivity":
            result = data.get("sensitivity", data)
            names = result.get("names", result.get("parameters"))
            mu = result.get("mu_star")
            sigma = result.get("sigma")
            if names is None or mu is None or sigma is None:
                raise ValueError("sensitivity figure needs names, mu_star, and sigma")
            ax.scatter(mu, sigma)
            for x, y, label in zip(mu, sigma, names):
                ax.annotate(label, (x, y), xytext=(4, 4), textcoords="offset points", fontsize=8)
            ax.set(xlabel="mean absolute elementary effect, μ*", ylabel="standard deviation, σ")
        elif kind == "projection":
            rows = data.get("projection", _rows(data))
            if isinstance(rows, dict):
                rows = rows.get("materials", rows.get("rows", []))
            distances = [row["distance"] for row in rows if row.get("distance") is not None]
            labels = [str(row.get("name", row.get("material", index))) for index, row in enumerate(rows)
                      if row.get("distance") is not None]
            if not distances:
                raise ValueError("projection figure needs measured-material distances")
            indices = np.argsort(distances)
            ax.barh(np.arange(len(indices)), np.asarray(distances)[indices])
            ax.set_yticks(np.arange(len(indices)), np.asarray(labels)[indices])
            ax.invert_yaxis()
            ax.set(xlabel=data.get("distance_label", "distance in recorded property coordinates"))
        else:
            rows = [row for row in _rows(data) if row.get("kind", "measured") in ("measured", "physical_zero")
                    and row.get("status", "complete") in ("complete", "done", "ok")]
            if not rows and any(name in data for name in ("mean", "estimate", "yield")):
                rows = [data]
            if not rows:
                raise ValueError("figure needs complete measured observations")
            recorded_units = {row["units"] for row in rows if row.get("units")}
            if len(recorded_units) > 1:
                raise ValueError("cannot compare responses with different recorded units")
            units = next(iter(recorded_units), data.get("units"))
            objective = data.get("objective", "")
            if units == "summed_junction_QPs_per_primary" or objective in ("total_qps_per_primary", "total_QPs_per_primary", "junction_qps_per_primary"):
                label = "summed junction QPs / primary"
            elif units == "weighted_junction_QPs_per_injected_eV" or objective in ("device_junction_qps_per_eV", "device_weighted_junction_qps_per_energy"):
                label = "weighted junction QPs / injected eV"
            elif units == "MHz us" or objective == "integrated_density":
                label = "integrated decoherence rate [MHz μs]"
            else:
                label = units.replace("_", " ") if units else "recorded response"
            label = data.get("yield_label", label)
            means = np.asarray([_estimate(row) for row in rows])
            errors = np.asarray([_uncertainty(row) for row in rows])
            if not np.isfinite(means).all() or not np.isfinite(errors).all() or (errors < 0).any():
                raise ValueError("invalid recorded means or uncertainties")
            if kind == "convergence":
                x = [row.get("n_sites", row.get("events", row.get("n", index + 1))) for index, row in enumerate(rows)]
                ax.errorbar(x, means, yerr=errors, marker="o", capsize=3)
                ax.set_xlabel(data.get("x_label", "sample count"))
            elif kind == "search":
                groups = sorted(set(str(row.get("method", "search")) for row in rows))
                for group in groups:
                    chosen = [index for index, row in enumerate(rows) if str(row.get("method", "search")) == group]
                    x = [rows[index].get("evaluation", order + 1) for order, index in enumerate(chosen)]
                    ax.plot(x, np.minimum.accumulate(means[chosen]), marker=".", label=group)
                ax.set_xlabel("complete candidate evaluations")
                ax.legend(frameon=False)
            else:
                indices = np.argsort(means)
                labels = [str(row.get("name", row.get("material", row.get("candidate_id", index)))) for index, row in enumerate(rows)]
                ax.errorbar(means[indices], np.arange(len(indices)), xerr=errors[indices], fmt="o", capsize=3)
                ax.set_yticks(np.arange(len(indices)), np.asarray(labels)[indices], fontsize=8)
                ax.invert_yaxis()
                ax.set_xlabel(label)
            if kind in ("convergence", "search"):
                ax.set_ylabel(label)
        ax.set_title(title or data.get("title", kind.replace("_", " ").capitalize()))
        ax.grid(alpha=.2)
        target = Path(output)
        if target.suffix.lower() not in (".png", ".pdf", ".svg"):
            raise ValueError("figure output must be PNG, PDF, or SVG")
        target.parent.mkdir(parents=True, exist_ok=True)
        fig.savefig(target, dpi=dpi)
    finally:
        plt.close(fig)
    return target


def export_presentation(figures, output, title="Material-property study"):
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as error:
        raise RuntimeError("presentation export needs python-pptx") from error
    presentation = Presentation()
    for index, path in enumerate(figures):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title if len(figures) == 1 else f"{title}: {index + 1}"
        slide.shapes.add_picture(str(path), Inches(.3), Inches(1), width=Inches(9.3))
    target = Path(output)
    target.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(target)
    return target
